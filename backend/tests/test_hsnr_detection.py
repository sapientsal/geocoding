"""Test HSNR column detection fix (Wolfsburg format)."""
import io
import os
import time
from pathlib import Path

import openpyxl
import pytest
import requests
from dotenv import dotenv_values
from pymongo import MongoClient

frontend_env = dotenv_values("/app/frontend/.env")
backend_env = dotenv_values("/app/backend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")).rstrip("/")
MONGO_URL = backend_env.get("MONGO_URL")
DB_NAME = backend_env.get("DB_NAME")


def build_wolfsburg_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["STRASSE", "HSNR", "HSNR_ZUSATZ", "PLZ", "ORT", "WE", "Restpoti", "WH1"])
    rows = [
        ("Am Finkenhaus", 1, "", 38444, "Wolfsburg", 2, 100, "OK"),
        ("Am Finkenhaus", 3, "", 38444, "Wolfsburg", 1, 50, "NEW"),
        ("Am Finkenhaus", 4, "a", 38444, "Wolfsburg", 3, 75, "OK"),
        ("Hauptstraße", 2, "", 38444, "Wolfsburg", 4, 200, "OK"),
        ("Hauptstraße", 4, "", 38444, "Wolfsburg", 1, 25, "NEW"),
    ]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture(scope="module")
def xlsx_bytes():
    return build_wolfsburg_xlsx().getvalue()


@pytest.fixture(scope="module")
def mongo_ids():
    return {"job_ids": []}


def test_1_preview_detects_hsnr(xlsx_bytes):
    files = {"file": ("wolfsburg_test.xlsx", xlsx_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/preview", files=files, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["detected_format"] == "German (separate columns)", data
    dc = data["detected_columns"]
    assert dc["house_number"] == "HSNR", dc
    assert dc["zusatz"] == "HSNR_ZUSATZ", dc
    assert dc["street"] == "STRASSE"
    assert dc["plz"] == "PLZ"
    assert dc["ort"] == "ORT"
    assert len(data["preview_addresses"]) == 5, data["preview_addresses"]
    first = data["preview_addresses"][0]["address"]
    assert "Am Finkenhaus" in first and "1" in first and "38444" in first and "Wolfsburg" in first, first


def test_2_e2e_upload_and_export(xlsx_bytes, mongo_ids):
    files = {"file": ("wolfsburg_test.xlsx", xlsx_bytes,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/upload-optimized", files=files, timeout=30)
    assert r.status_code == 200, r.text
    job_id = r.json().get("job_id")
    assert job_id
    mongo_ids["job_ids"].append(job_id)

    # Poll job
    status = None
    job_data = None
    for _ in range(60):
        jr = requests.get(f"{BASE_URL}/api/job/{job_id}", timeout=15)
        assert jr.status_code == 200, jr.text
        job_data = jr.json()
        status = job_data.get("status")
        if status in ("completed", "failed", "error"):
            break
        time.sleep(2)

    assert status == "completed", f"Job status={status}, data={job_data}"
    geocoded = job_data.get("geocoded_count") or job_data.get("addresses_geocoded") or 0
    # fallback: check optimized result
    opt = requests.get(f"{BASE_URL}/api/optimized/{job_id}", timeout=15)
    assert opt.status_code == 200, opt.text
    opt_data = opt.json()
    addresses = opt_data.get("optimized_addresses") or opt_data.get("addresses") or opt_data.get("route") or []
    assert len(addresses) > 0, f"no geocoded addresses"
    geocoded_ok = [a for a in addresses if a.get("geocoded")]
    assert len(geocoded_ok) == 5, f"expected 5 geocoded, got {len(geocoded_ok)}"

    # Mongo persistence
    client = MongoClient(MONGO_URL)
    db = client[DB_NAME]
    route_doc = db["routes"].find_one({"job_id": job_id})
    assert route_doc is not None, "route doc missing in Mongo"
    # Check row_data preserves extra columns
    addrs = route_doc.get("optimized_addresses") or route_doc.get("addresses") or route_doc.get("optimized_route") or []
    found_extras = False
    for a in addrs:
        rd = a.get("row_data") or {}
        keys_upper = {k.upper() for k in rd.keys()}
        if "WE" in keys_upper and "RESTPOTI" in keys_upper and "WH1" in keys_upper:
            found_extras = True
            break
    assert found_extras, f"extra cols missing in row_data; sample={addrs[:1]}"
    client.close()

    # Export
    ex = requests.get(f"{BASE_URL}/api/optimized/{job_id}/export", timeout=30)
    assert ex.status_code == 200, ex.text
    assert "spreadsheetml" in ex.headers.get("content-type", "") or ex.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(ex.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    header_upper = [str(h).upper() if h else "" for h in header]
    assert "WH1" in header_upper, f"WH1 not in export header: {header}"
    assert "RESTPOTI" in header_upper, f"Restpoti not in export header: {header}"


def test_zzz_cleanup(mongo_ids):
    client = MongoClient(MONGO_URL)
    db = client[DB_NAME]
    for jid in mongo_ids["job_ids"]:
        db["upload_jobs"].delete_many({"job_id": jid})
        db["routes"].delete_many({"job_id": jid})
    client.close()
