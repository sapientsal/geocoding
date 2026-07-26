"""E2E tests for /api/upload-optimized: Haldensleben-format, Export, Legacy-format."""
import io
import os
import time
import pytest
import requests
import pandas as pd
from dotenv import dotenv_values
from pymongo import MongoClient
from openpyxl import load_workbook

frontend_env = dotenv_values("/app/frontend/.env")
BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or frontend_env.get("REACT_APP_BACKEND_URL")).rstrip("/")

backend_env = dotenv_values("/app/backend/.env")
MONGO_URL = backend_env.get("MONGO_URL")
DB_NAME = backend_env.get("DB_NAME")

# Track test job ids for cleanup
_test_job_ids = []


@pytest.fixture(scope="module")
def db():
    client = MongoClient(MONGO_URL)
    yield client[DB_NAME]
    # cleanup
    for jid in _test_job_ids:
        client[DB_NAME]["upload_jobs"].delete_many({"id": jid})
        client[DB_NAME]["routes"].delete_many({"job_id": jid})
    client.close()


def _upload_df(df: pd.DataFrame, filename: str) -> str:
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine="openpyxl")
    buf.seek(0)
    r = requests.post(
        f"{BASE_URL}/api/upload-optimized",
        files={"file": (filename, buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        timeout=30,
    )
    assert r.status_code == 200, f"upload failed: {r.status_code} {r.text[:300]}"
    jid = r.json()["job_id"]
    _test_job_ids.append(jid)
    return jid


def _wait_for_completion(job_id: str, timeout_s: int = 180) -> dict:
    start = time.time()
    last_status = None
    while time.time() - start < timeout_s:
        r = requests.get(f"{BASE_URL}/api/job/{job_id}", timeout=15)
        assert r.status_code == 200, f"job status HTTP {r.status_code}"
        data = r.json()
        if data["status"] != last_status:
            last_status = data["status"]
            print(f"[{job_id[:8]}] status={data['status']} processed={data.get('processed_addresses')}/{data.get('total_addresses')}")
        if data["status"] == "completed":
            return data
        if data["status"] == "error":
            pytest.fail(f"Job errored: {data.get('error_message')}")
        time.sleep(2)
    pytest.fail(f"Job {job_id} did not complete within {timeout_s}s (last status={last_status})")


class TestHaldenslebenE2E:
    """Test 1: Upload Haldensleben-format Excel, verify hierarchical sort."""

    def test_haldensleben_upload_and_sort(self, db):
        df = pd.DataFrame({
            'U_ASSIGNED_ISP': ['O2'] * 6,
            'STREET_NAME': ['Lange Straße', 'Lange Straße', 'Lange Straße', 'Lange Straße', 'Bahnhofstraße', 'Bahnhofstraße'],
            'NUMBER': [1, 2, 1, 2, 3, 5],
            'NUMBER_AFFIX': [None, None, None, None, None, None],
            'DISTRICT': ['Uthmöden', 'Uthmöden', 'Haldensleben', 'Haldensleben', 'Haldensleben', 'Haldensleben'],
            'PLZ': [39345, 39345, 39340, 39340, 39340, 39340],
            'MUNICIPALITY': ['Haldensleben'] * 6,
        })
        job_id = _upload_df(df, "TEST_haldensleben.xlsx")
        job = _wait_for_completion(job_id, timeout_s=240)
        assert job["total_addresses"] == 6

        route = db["routes"].find_one({"job_id": job_id})
        assert route is not None, "No route document found"
        opts = route["optimized_addresses"]
        assert len(opts) == 6, f"expected 6 optimized addresses got {len(opts)}"

        districts = [o["row_data"].get("DISTRICT") for o in opts]
        streets = [o["row_data"].get("STREET_NAME") for o in opts]
        numbers = [o["row_data"].get("NUMBER") for o in opts]
        print("Order:", list(zip(districts, streets, numbers)))

        # (a) Uthmöden block contiguous and before Haldensleben
        segs, prev = 0, None
        for d in districts:
            if d != prev:
                segs += 1
                prev = d
        assert segs == 2, f"District blocks not contiguous, got {segs} segments: {districts}"
        assert districts[0] == 'Uthmöden', f"Northernmost Uthmöden must start, got {districts[0]}"
        assert districts.count('Uthmöden') == 2 and districts.count('Haldensleben') == 4

        # (b) Lange Straße of Uthmöden not interleaved with Haldensleben Lange Straße
        uth_streets = [s for s, d in zip(streets, districts) if d == 'Uthmöden']
        assert all(s == 'Lange Straße' for s in uth_streets)

        # (c) house numbers ascending within each (district, street) group
        from itertools import groupby
        groups = [(k, [n for _, _, n in g]) for k, g in groupby(zip(districts, streets, numbers), key=lambda x: (x[0], x[1]))]
        for k, nums in groups:
            assert nums == sorted(nums), f"house numbers not ascending in {k}: {nums}"

        # (d) street_city_key includes Ortsteil (either OT name or plausibly the composite)
        for o in opts:
            key = o["row_data"].get("street_city_key")
            assert key is not None, "street_city_key missing"
            ot = o["row_data"].get("DISTRICT")
            assert ot in key, f"Ortsteil '{ot}' not in street_city_key '{key}'"

        # Persist job_id for export test via class attribute
        TestHaldenslebenE2E.job_id = job_id

    def test_export_haldensleben(self):
        job_id = getattr(TestHaldenslebenE2E, "job_id", None)
        assert job_id, "no job_id from previous test"
        r = requests.get(f"{BASE_URL}/api/optimized/{job_id}/export", timeout=60)
        assert r.status_code == 200, f"export HTTP {r.status_code}: {r.text[:200]}"
        assert len(r.content) > 100
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        for required in ['STREET_NAME', 'NUMBER', 'DISTRICT', 'PLZ', 'MUNICIPALITY']:
            assert required in headers, f"missing column {required} in export headers: {headers}"
        # rows count
        row_count = ws.max_row - 1
        assert row_count == 6, f"expected 6 data rows got {row_count}"


class TestLegacyFormatE2E:
    """Test 3: Legacy German columns Straße/Hausnummer/PLZ/Ort still works."""

    def test_legacy_upload(self, db):
        df = pd.DataFrame({
            'Straße': ['Hauptstraße'] * 4,
            'Hausnummer': [1, 2, 3, 4],
            'PLZ': [27726] * 4,
            'Ort': ['Worpswede'] * 4,
        })
        job_id = _upload_df(df, "TEST_legacy_worpswede.xlsx")
        job = _wait_for_completion(job_id, timeout_s=180)
        assert job["total_addresses"] == 4
        assert job["status"] == "completed"

        # Export must work
        r = requests.get(f"{BASE_URL}/api/optimized/{job_id}/export", timeout=60)
        assert r.status_code == 200, f"legacy export HTTP {r.status_code}: {r.text[:200]}"
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        for required in ['Straße', 'Hausnummer', 'PLZ', 'Ort']:
            assert required in headers, f"missing column {required} in export: {headers}"
